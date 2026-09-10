# -*- coding: utf-8 -*-
"""
law_tool_app.py — 法規鑑別追蹤工具

本機用法（雙擊 start_law_tool.bat，或手動執行）：
    pip install -r requirements.txt
    python law_tool_app.py

執行後會自動開啟瀏覽器，網址是 http://127.0.0.1:5050 。
資料存在同一個資料夾裡的 law_tool.db 檔案，關掉程式、重開電腦都不會不見。
這支程式必須在「有一般網路連線」的電腦上執行（你自己的電腦即可），不能在
雲端沙盒環境裡執行。

也可以部署到雲端主機（例如 Railway），變成一個固定網址、任何電腦開瀏覽器
就能用，不用再依賴某一台電腦。部署時用 gunicorn 啟動（見 Procfile），並且
一定要接上「永久儲存空間（Volume）」、把 LAW_TOOL_DB_PATH 指向那個路徑，
否則服務重啟時資料庫會被清空。詳細步驟見 README.md。

這是 moj_law_diff.py 的延伸：把「抓最新條文、比對舊版本、跑文字 diff」這套
已經驗證過的邏輯，包成一個可以長期使用的本機網頁系統，取代原本要一條一條
法規手動下指令的方式，並且把每一條的鑑別結果（適用性、守規性評估、鑑別人員、
是否結案）存起來，下次可以繼續編輯，也可以隨時匯出成 Excel。

功能：
    - 新增要追蹤的法規（給法規代碼 pcode，並設定「追蹤起始日」）
    - 「立即比對」：抓最新條文，跟上次確認到的日期之後比對，只列出新的異動
    - 在網頁上直接填寫條文適用性／守規性評估／鑑別人員／備註／是否結案
    - 匯出單一法規或全部法規的 Excel，欄位對齊你原本的法規鑑別表格式，
      有異動的文字自動用粗體＋底線標示
"""

import json
import os
import re
import sqlite3
import sys
import threading
import time
import warnings
import webbrowser
from datetime import date, datetime, timedelta

import requests
from bs4 import BeautifulSoup
from flask import Flask, flash, redirect, render_template_string, request, send_file, url_for
from markupsafe import Markup, escape

# ---------------------------------------------------------------------------
# 勞動部勞動法令查詢系統 抓取／比對引擎（laws.mol.gov.tw）
#
# 原本這支工具是打「全國法規資料庫」（law.moj.gov.tw），改成打這個勞動部
# 自己的系統，主要是因為對職業安全衛生法這種法規，勞動法令查詢系統的沿革
# 頁面「每一次修正」都有留舊版全文可以直接抓（法規沿革頁面每一列都附
# 「所有條文」連結），全國法規資料庫反而常常只留現在這一版，逼得工具要用
# 公告文字反推異動條號、猜不到舊條文內容。這裡换成用舊版全文做逐字比對，
# 準確度好很多，公告文字解析（parse_history_entries）留著當「萬一某部法規
# 也一樣沒留舊版全文」時的備援，邏輯完全比照原本那一版。
# ---------------------------------------------------------------------------

BASE = "https://laws.mol.gov.tw"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (compatible; RegComplianceTool/0.2; internal use)"
}
ARTICLE_NO_RE = re.compile(r"第\s*[\d\-]+\s*條")
_WARNED_INSECURE = False


def http_get(url: str):
    """對勞動部勞動法令查詢系統發出 GET 請求；遇到某些 Windows 環境常見的
    憑證驗證相容性錯誤時，自動退回不驗證憑證的方式重試一次。"""
    global _WARNED_INSECURE
    try:
        resp = requests.get(url, headers=HEADERS, timeout=20)
    except requests.exceptions.SSLError:
        if not _WARNED_INSECURE:
            print(
                "提醒：本機的憑證驗證對 laws.mol.gov.tw 回報相容性錯誤，"
                "已自動改用不驗證憑證的方式重新連線（僅讀取公開政府法規網站，風險低）。",
                file=sys.stderr,
            )
            _WARNED_INSECURE = True
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            resp = requests.get(url, headers=HEADERS, timeout=20, verify=False)
    resp.raise_for_status()
    return resp


def normalize_article_no(raw: str) -> str:
    m = ARTICLE_NO_RE.search(raw)
    text = m.group(0) if m else raw
    return re.sub(r"\s+", "", text)


def fetch_articles(url: str) -> dict:
    """抓某一頁『所有條文』（現行版 FLAWDAT0201 或指定日期的舊版 FLAWDAT08）
    的內容，回傳 {正規化條號: 條文內容}。"""
    resp = http_get(url)
    resp.encoding = resp.apparent_encoding or "utf-8"
    soup = BeautifulSoup(resp.text, "html.parser")
    articles = {}
    for row in soup.select(".row"):
        no_el = row.select_one(".col-no")
        data_el = row.select_one(".col-data")
        if not no_el or not data_el:
            continue
        raw_no = no_el.get_text(strip=True)
        if "條" not in raw_no:
            continue
        no = normalize_article_no(raw_no)
        text = data_el.get_text("\n", strip=True)
        text = re.sub(r"\n{2,}", "\n", text)
        articles[no] = text
    return articles


def roc_to_yyyymmdd(s: str):
    """把「民國114年12月19日」這種阿拉伯數字寫法的日期轉成 yyyymmdd。法規
    摘要欄位（公(發)布日期／修正日期）用的是這種格式；沿革逐條公告文字用的
    是純國字數字，要用 roc_cn_to_yyyymmdd()。"""
    if not s:
        return None
    m = re.search(r"(\d+)\s*年\s*(\d+)\s*月\s*(\d+)\s*日", s)
    if not m:
        return None
    roc_year, month, day = (int(x) for x in m.groups())
    return f"{roc_year + 1911:04d}{month:02d}{day:02d}"


def yyyymmdd_to_roc_display(d: str) -> str:
    y, mo, da = int(d[0:4]), int(d[4:6]), int(d[6:8])
    return f"民國 {y - 1911} 年 {mo:02d} 月 {da:02d} 日"


def fmt_date(d: str) -> str:
    if not d or len(d) != 8:
        return d or ""
    return f"{d[0:4]}-{d[4:6]}-{d[6:8]}"


# ---------------------------------------------------------------------------
# 沿革公告文字解析：勞動法令查詢系統對每一次修正都會留一段公告原文（例如
# 「修正公布第2、6條條文；增訂第15-1條條文」），這裡解析成結構化的「這次
# 修正/增訂/刪除了哪些條號」，並且額外算出「這次修正實際生效的日期」——
# 很多修正是「公告日」跟「施行日」不同天（甚至公告時寫「施行日期，由行政院
# 定之」，實際日期要等後續另一則公告才知道），你特別交代過要排除「已公告
# 但還沒生效」的修正，所以這裡不能只看公告日期，要盡量抓出真正生效的那天，
# 抓不到明確生效日時保守當作跟公告日同一天生效。
#
# 這段解析的對象是政府公告的自然語言文字，用詞、標點不完全一致（有的用
# 「修正公布」有的用「修正發布」，「－」跟「～」意義也不同：「～」是「從～到～」
# 的條號範圍，「－」是「之幾」的子條號如「第15-1條」），已經涵蓋常見的幾種
# 寫法，但無法保證每一種罕見寫法都解析得出來；解析不出具體條號時，呼叫端
# 一律會退回「請自行到沿革頁面確認」的警示，不會因為解析失敗就誤報成
# 「沒有異動」。
# ---------------------------------------------------------------------------

_CN_NUM_MAP = {"零": 0, "〇": 0, "一": 1, "二": 2, "兩": 2, "三": 3, "四": 4,
               "五": 5, "六": 6, "七": 7, "八": 8, "九": 9}
_CN_UNIT_MAP = {"十": 10, "百": 100, "千": 1000}
_CN_NUM_CLASS = "[一二三四五六七八九十百千零〇兩]"


def cn_num_to_int(s: str):
    """把中文數字（如「一百十四」「九十一」「十九」）轉成阿拉伯數字，用來讀
    沿革公告裡「中華民國一百十四年十二月十九日」這種完全用國字寫的日期
    （法規本身摘要欄位的日期是阿拉伯數字，但沿革逐條公告文字習慣用國字，
    兩邊要分開處理）。只處理到千位，日期用途绰绰有餘。"""
    if not s:
        return None
    result, pending = 0, 0
    if s[0] == "十":
        result = 10
        s = s[1:]
    for ch in s:
        if ch in _CN_NUM_MAP:
            pending = _CN_NUM_MAP[ch]
        elif ch in _CN_UNIT_MAP:
            result += (pending if pending else 1) * _CN_UNIT_MAP[ch]
            pending = 0
    return result + pending


def roc_cn_to_yyyymmdd(text: str):
    m = re.search(rf"中華民國({_CN_NUM_CLASS}+)年({_CN_NUM_CLASS}+)月({_CN_NUM_CLASS}+)日", text)
    if not m:
        return None
    y, mo, d = (cn_num_to_int(g) for g in m.groups())
    if y is None or mo is None or d is None:
        return None
    return f"{y + 1911:04d}{mo:02d}{d:02d}"


def _entry_effective_date(compact_text: str, promulgation_date: str) -> str:
    """從沿革單筆公告文字判斷『這次修正真正生效』的日期。公告文字常見寫法：
    「自即日施行」「自發布日施行」＝跟公告日同一天；「自OO年OO月OO日施行」＝
    明確指定一天；有時同一則公告會分好幾批生效（不同條文分批），這裡採保守
    做法——抓到的候選生效日一律取『最晚』的一個，寧可晚一點才把這次修正當
    作已生效，也不要把還沒全部生效的修正提早當成現行條文（這是你特別要求
    的：預告中、還沒生效的一律不算）。完全找不到生效日說明的罕見情況，保守
    當作跟公告日同一天生效。"""
    dates_found = []
    if re.search(r"自(即日|發布日|公布日)施行", compact_text):
        dates_found.append(promulgation_date)
    for m in re.finditer(rf"自({_CN_NUM_CLASS}+)年({_CN_NUM_CLASS}+)月({_CN_NUM_CLASS}+)日施行", compact_text):
        y, mo, d = (cn_num_to_int(g) for g in m.groups())
        if y is not None and mo is not None and d is not None:
            dates_found.append(f"{y + 1911:04d}{mo:02d}{d:02d}")
    if not dates_found:
        return promulgation_date
    return max(dates_found)


def _expand_article_tokens(numlist: str) -> list:
    """把「2、6、9、43～46」這種用頓號分隔、可能帶「～」範圍或「－」子條號的
    條號清單，展開成正規化的條號字串列表（跟 normalize_article_no() 產生的
    格式一致，例如「第43條」「第15-1條」），才能跟抓下來的條文內容對上。"""
    tokens = [t for t in re.split(r"[、，]", numlist) if t]
    out = []
    for tok in tokens:
        m = re.match(r"^(\d+)-(\d+)～(\d+)-(\d+)$", tok)
        if m:
            base1, a, base2, c = m.groups()
            if base1 == base2 and int(a) <= int(c):
                out.extend(f"第{base1}-{i}條" for i in range(int(a), int(c) + 1))
            else:
                out.extend([f"第{base1}-{a}條", f"第{base2}-{c}條"])
            continue
        m = re.match(r"^(\d+)～(\d+)$", tok)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            if a <= b:
                out.extend(f"第{i}條" for i in range(a, b + 1))
            else:
                out.append(f"第{tok}條")
            continue
        out.append(f"第{tok}條")
    return out


def parse_history_entries(soup) -> list:
    """解析法規沿革頁面（.law-history）裡每一次公告的說明文字，回傳每次公告
    的公告日期（date）、實際生效日期（effective_date），以及這次「修正」
    「增訂」「刪除」了哪些條號（找不到具體條號、或這次是修正/制定「全文」
    沒有逐條列出的，對應欄位就是空 list／full_revision=True）。is_latest 標記
    是不是沿革清單裡最新的那一筆——勞動法令查詢系統對「目前最新、還沒被
    取代」的版本，全文要去現行條文頁（FLAWDAT0201）抓，其餘已被取代的舊
    版本才能用歷史條文頁（FLAWDAT08+ldate）抓，這個判斷只看『是不是清單
    裡最新一筆』，跟它是否已經生效無關（即使最新這筆還沒生效，網站上它也
    還不是「歷史版本」）。
    依公告日期新到舊排序，跟 list_versions() 的 versions 排序方式一致。"""
    entries = []
    rows = soup.select(".law-history .row")
    for i, row in enumerate(rows):
        data_el = row.select_one(".col-data")
        if not data_el:
            continue
        raw = data_el.get_text("", strip=True)
        compact = re.sub(r"\s+", "", raw)
        if not re.match(r"^\d+\.", compact):
            continue
        entry_date = roc_cn_to_yyyymmdd(compact)
        if not entry_date:
            continue

        full_revision = bool(re.search(r"(修正|制定|訂定)(公布|發布)(名稱及)?全文\d+條", compact))

        amended, added, deleted = [], [], []
        for verb, nums in re.findall(r"(修正公布|修正發布|增訂|刪除)第([\d、，\-～]+)條", compact):
            arts = _expand_article_tokens(nums)
            if verb in ("修正公布", "修正發布"):
                amended.extend(arts)
            elif verb == "增訂":
                added.extend(arts)
            elif verb == "刪除":
                deleted.extend(arts)

        entries.append({
            "date": entry_date,
            "effective_date": _entry_effective_date(compact, entry_date),
            "raw_text": raw, "full_revision": full_revision,
            "amended": amended, "added": added, "deleted": deleted,
        })
    # 頁面上的沿革列表是按公告時間「舊到新」排列（第一列是最早的制定公布），
    # 不是新到舊，所以「是不是最新一筆」不能看 DOM 原始 index，要照公告日期
    # 排序後才知道；排序完 entries[0] 才是真正最新的那一筆。
    entries.sort(key=lambda e: e["date"], reverse=True)
    for i, e in enumerate(entries):
        e["is_latest"] = i == 0
    return entries


def _shift_day(d: str, delta: int) -> str:
    dt_obj = date(int(d[0:4]), int(d[4:6]), int(d[6:8])) + timedelta(days=delta)
    return dt_obj.strftime("%Y%m%d")


def previous_day(d: str) -> str:
    return _shift_day(d, -1)


def next_day(d: str) -> str:
    return _shift_day(d, 1)


def list_versions(pcode: str) -> dict:
    """pcode 在這個資料來源其實是勞動法令查詢系統的法規代碼 id（例如
    「FL015013」，格式跟舊的全國法規資料庫 pcode 不一樣，但沿用同一個欄位
    名稱、同一套資料庫結構，改動範圍降到最小）。"""
    url = f"{BASE}/FLAW/FLAWDAT07.aspx?id={pcode}"
    resp = http_get(url)
    resp.encoding = resp.apparent_encoding or "utf-8"
    soup = BeautifulSoup(resp.text, "html.parser")

    name_el = soup.find(string=re.compile("法規名稱"))
    law_name = None
    if name_el:
        row = name_el.find_parent("tr") or name_el.find_parent()
        if row:
            law_name = row.get_text(" ", strip=True).replace("法規名稱：", "").strip()
            law_name = re.sub(r"\s*英\s*$", "", law_name).strip()

    amend_date_display = None
    for label in ("修正日期", "公(發)布日期", "發布日期"):
        el = soup.find(string=re.compile(re.escape(label)))
        if not el:
            continue
        row = el.find_parent("tr") or el.find_parent()
        if row:
            amend_date_display = row.get_text(" ", strip=True).replace(f"{label}：", "").strip()
            break

    history_entries = parse_history_entries(soup)

    versions = [
        {"date": e["effective_date"], "ldate": e["date"], "source": "current" if e["is_latest"] else "old"}
        for e in history_entries
    ]
    versions.sort(key=lambda v: v["date"], reverse=True)

    # 如果沿革清單裡最新一筆的生效日還沒到（今天還沒到那一天），首頁摘要的
    # 「現行修正日期」不能直接顯示網站原本標的那個修正日期，不然會讓人誤以
    # 為那次修正已經生效——改成顯示『目前真正有效』那個版本（也就是這次還
    # 沒生效之前，實際還在適用的上一版）的公告日期。
    today_str = date.today().strftime("%Y%m%d")
    still_effective = [v for v in versions if v["date"] <= today_str]
    if still_effective and history_entries and still_effective[0]["ldate"] != history_entries[0]["date"]:
        eff_entry = next((e for e in history_entries if e["date"] == still_effective[0]["ldate"]), None)
        if eff_entry:
            amend_date_display = yyyymmdd_to_roc_display(eff_entry["date"])

    return {
        "law_name": law_name, "amend_date_display": amend_date_display,
        "versions": versions, "history_entries": history_entries,
    }


def search_laws_by_keyword(keyword: str, max_pages: int = 2) -> tuple:
    """用法規名稱關鍵字查詢勞動部勞動法令查詢系統，回傳可能對應的法規清單。

    一開始是直接打整合查詢（不帶 type 參數），結果會把關鍵字出現在『條文
    內容』裡、甚至其他分類（行政規則、解釋令函）的東西全部混在一起、依
    相關度排序，導致「職業安全衛生法」這種關鍵字搜尋出一堆名稱完全不相關
    的法規（因為條文裡剛好提到這幾個字）——這裡改成帶 type=name,01,02,03
    這個查詢參數，只比對『法規名稱』本身有沒有出現關鍵字（網站上「法規
    查詢／法規名稱」那個篩選分頁用的就是這個參數），結果精確很多，也已經
    是部分比對（不用打完整全名一樣找得到）。這個篩選本身就不包含「法規
    草案」（草案是另一個獨立的 type=drafts 分類），天然符合『只採用已經
    生效的法規，草案不算』的要求，不用再另外判斷類別或生效狀態。
    這個篩選底下的結果如果是已廢止的法規，名稱前面會帶一個 class="fei"
    的「廢」字樣提示，這裡順便解析出來存進 repealed 欄位，畫面上會顯示
    「已廢止」提示，由使用者自己決定要不要繼續追蹤。
    第二個回傳值代表結果是否被截斷（代表還有更多筆，建議使用者輸入更精確
    的關鍵字）。"""
    from urllib.parse import quote

    results = []
    truncated = False
    for page in range(1, max_pages + 1):
        url = f"{BASE}/results.aspx?searchmode=global&keyword={quote(keyword)}&type=name,01,02,03&page={page}"
        resp = http_get(url)
        resp.encoding = resp.apparent_encoding or "utf-8"
        soup = BeautifulSoup(resp.text, "html.parser")
        table = soup.select_one("table.laws-table") or soup.select_one("table")
        if not table:
            break
        rows = table.select("tr")[1:]  # 跳過表頭那一列（序／附件或筆數／法規名稱／異動日期）
        if not rows:
            break
        for tr in rows:
            tds = tr.find_all("td")
            if len(tds) < 4:
                continue
            a = tds[2].find("a")
            if not a:
                continue
            m = re.search(r"[?&]id=([A-Za-z0-9]+)", a.get("href", ""))
            if not m:
                continue
            fei_span = a.find("span", class_="fei")
            repealed = fei_span is not None
            if fei_span:
                fei_span.extract()  # 拿掉「廢」字樣，才不會混進法規名稱裡
            results.append({
                "pcode": m.group(1),
                "law_name": a.get_text(strip=True),
                "amend_date_display": tds[3].get_text(strip=True),
                "repealed": repealed,
            })
        if len(rows) < 20:
            break
        if page == max_pages:
            truncated = True
    return results, truncated


def version_asof(versions_desc: list, target_date: str, strict_before: bool = False):
    for v in versions_desc:
        if (v["date"] < target_date) if strict_before else (v["date"] <= target_date):
            return v
    return None


def fetch_version_articles(pcode: str, version: dict) -> dict:
    """version["date"] 存的是『生效日』（比對邏輯用這個），但抓舊版全文網址
    要用『公告日』（version["ldate"]）——勞動法令查詢系統的歷史條文頁是用
    公告日當查詢參數，這兩個日期在「公告後延後施行」的修正裡並不是同一天。
    is_latest（也就是 source=="current"）的版本要用現行條文頁抓，不能用
    歷史條文頁：這個網站對「還沒被取代的最新版本」帶公告日查歷史條文頁會
    直接回錯誤頁。"""
    if version["source"] == "current":
        return fetch_articles(f"{BASE}/FLAW/FLAWDAT0201.aspx?id={pcode}")
    return fetch_articles(f"{BASE}/FLAW/FLAWDAT08.aspx?id={pcode}&ldate={version['ldate']}")


def inline_diff(old_text: str, new_text: str):
    import difflib

    sm = difflib.SequenceMatcher(None, old_text, new_text, autojunk=False)
    old_runs, new_runs = [], []
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag == "equal":
            old_runs.append((old_text[i1:i2], False))
            new_runs.append((new_text[j1:j2], False))
        else:
            if i1 != i2:
                old_runs.append((old_text[i1:i2], True))
            if j1 != j2:
                new_runs.append((new_text[j1:j2], True))
    return old_runs, new_runs


def _history_based_rows(pcode: str, history_entries: list, since: str, until: str, new_v: dict) -> tuple:
    """在完全抓不到舊版全文可以自動比對時（baseline_missing）的備援方案：改
    用沿革頁面的公告文字，找出「追蹤起始日之後」公告過的那幾次修正，從公告
    原文解析出實際異動了哪些條號，逐條列進待鑑別清單——新增的條文可以直接
    抓現行條文全文當作『新增後的內容』（新增條文本來就沒有舊條文可比，這點
    跟平常真的能比對到舊版本時的處理方式一致）；修正/刪除的條文因為真的沒有
    舊條文可比，只能標注清楚、附上沿革頁面連結，請使用者自行核對。
    回傳 (changed_rows, extra_note)：extra_note 是要額外附加在 note 欄位、
    講清楚這份清單是怎麼來的一句話；解析不到任何具體條號時 changed_rows 只有
    一筆通用警示（跟完全沒有沿革資料可解析時的後備行為一樣），確保不會因為
    公告文字解析失敗就悄悄漏掉這次修正。"""
    qualifying = [e for e in history_entries if since and since < e["date"] <= until]
    hist_url = f"{BASE}/FLAW/FLAWDAT07.aspx?id={pcode}"

    def generic_warning():
        warn_text = (
            f"勞動法令查詢系統沒有保留 {fmt_date(new_v['date'])} 這次修正之前的完整條文"
            f"（該法規沿革頁面沒有提供舊版全文連結），系統無法自動抓出實際異動了哪些條文。"
            f"這次修正的日期在追蹤起始日之後，請務必到勞動法令查詢系統的「法規沿革」頁面"
            f"（{hist_url}）查看公告內容，自行確認異動條文並完成鑑別，鑑別完再把這一筆勾選結案。"
        )
        return [{
            "article_no": "（系統無法自動比對）", "status": "警示",
            "old_runs": [(warn_text, True)],
            "new_runs": [(f"（請自行查閱 {fmt_date(new_v['date'])} 生效的現行條文全文）", False)],
        }]

    if not qualifying:
        return generic_warning(), None

    if any(e["full_revision"] for e in qualifying):
        dates = "、".join(fmt_date(e["date"]) for e in qualifying if e["full_revision"])
        warn_text = (
            f"{dates} 的公告是修正／制定「全文」，公告原文沒有逐條列出異動條號，"
            f"勞動法令查詢系統也沒有保留修正前的舊條文，系統無法自動比對出差異，"
            f"請至沿革頁面（{hist_url}）查看公告內容，並將全部條文都列入這次鑑別範圍。"
        )
        rows = [{
            "article_no": "（全文修正，需整部重新鑑別）", "status": "警示",
            "old_runs": [(warn_text, True)],
            "new_runs": [(f"（請自行查閱 {fmt_date(new_v['date'])} 生效的現行條文全文）", False)],
        }]
        return rows, "偵測到全文修正／制定公告，公告文字沒有逐條列出條號，已提醒需整部重新鑑別。"

    current = fetch_version_articles(pcode, new_v)
    rows, seen = [], set()
    for e in qualifying:
        for no in e["added"]:
            if no in seen:
                continue
            seen.add(no)
            text = current.get(no) or "（找不到現行條文內容，請自行至法規內容頁查閱）"
            rows.append({"article_no": no, "status": "新增", "old_runs": [], "new_runs": [(text, True)]})
        for no in e["amended"]:
            if no in seen:
                continue
            seen.add(no)
            text = current.get(no) or "（找不到現行條文內容，請自行至法規內容頁查閱）"
            old_text = (
                f"勞動法令查詢系統未保留 {fmt_date(e['date'])} 這次修正前的舊條文，"
                f"請至沿革頁面（{hist_url}）查看公告內容自行比對修正前後差異。"
            )
            rows.append({"article_no": no, "status": "修正", "old_runs": [(old_text, True)], "new_runs": [(text, True)]})
        for no in e["deleted"]:
            if no in seen:
                continue
            seen.add(no)
            old_text = (
                f"此條文已於 {fmt_date(e['date'])} 被刪除，勞動法令查詢系統未保留刪除前的條文全文，"
                f"請至沿革頁面（{hist_url}）查看公告內容。"
            )
            rows.append({"article_no": no, "status": "刪除", "old_runs": [(old_text, True)], "new_runs": []})

    if not rows:
        return generic_warning(), None

    entry_dates = "、".join(fmt_date(e["date"]) for e in qualifying)
    return rows, (
        f"以上是根據沿革頁面 {entry_dates} 公告文字解析出來的異動條號（不是逐字比對舊條文），"
        f"「修正」「刪除」的條文因為資料庫沒有保留舊條文，請自行到沿革頁面核對修正前的內容。"
    )


def build_report(pcode: str, since: str = None, until: str = None):
    """比對某一部法規的新舊條文；given since，抓『since 之前最後生效的版本』
    對『until（預設今天）當時生效的版本』；沒給 since 就跟現行版本對上一版比。"""
    info = list_versions(pcode)
    law_name = info["law_name"]
    if not law_name:
        raise RuntimeError(f"抓不到 pcode={pcode} 的法規名稱，請確認代碼是否正確。")

    versions = info["versions"]
    amend_date = info["amend_date_display"]
    baseline_missing = False

    if since is None:
        if len(versions) < 2:
            return {
                "pcode": pcode, "law_name": law_name, "amend_date": amend_date,
                "old_version_date": None, "new_version_date": None,
                "amendment_dates_in_range": None, "note": None, "changed_rows": [],
                "baseline_missing": False,
            }
        new_v, old_v, note = versions[0], versions[1], None
        until = until or versions[0]["date"]
    else:
        until = until or date.today().strftime("%Y%m%d")
        old_v = version_asof(versions, since, strict_before=True)
        new_v = version_asof(versions, until, strict_before=False)
        note = None
        if new_v is None:
            raise RuntimeError(f"{fmt_date(until)} 這個日期比這部法規最早的資料還早，查不到任何版本。")
        if old_v is None:
            old_v = versions[-1]
            baseline_missing = True
            note = (
                f"這個資料庫對此法規最早只能追溯到 {fmt_date(old_v['date'])}，"
                f"更早之前的版本無法比對，已改用這個最早版本當基準。"
            )

    if old_v["date"] == new_v["date"]:
        changed_rows = []
        # baseline_missing 代表：往前找不到「追蹤起始日之前」有效的版本，只好
        # 拿資料庫裡找得到的最早版本（old_v）當基準——但這裡它又剛好跟現在
        # 生效的版本（new_v）是同一份，等於完全沒有更早的舊條文可以比對。
        # 這時候絕對不能直接回報「沒有新的修正」：現行版本的修正日期
        # （new_v）如果就落在追蹤起始日之後，代表追蹤期間內其實真的發生過
        # 修正，只是勞動法令查詢系統沒有保留修正前的舊條文、系統沒辦法自動逐字
        # 比對出異動了哪些條文而已。這時候改用沿革頁面的公告文字當備援資料
        # 來源：從公告原文解析出這次實際修正/增訂/刪除了哪些條號，把每一條
        # 分開列進待鑑別清單（新增的條文可以直接附上現行條文全文），逼使用者
        # 針對「修正」「刪除」的條文自己到沿革頁面核對前後差異，鑑別完再手動
        # 勾選結案——不管解析成不成功，都不會像以前一樣悄悄回報「沒有異動」
        # 而漏掉這次修正。
        if baseline_missing:
            changed_rows, extra_note = _history_based_rows(
                pcode, info.get("history_entries") or [], since, until, new_v
            )
            if extra_note:
                note = f"{note}{extra_note}"
        return {
            "pcode": pcode, "law_name": law_name, "amend_date": amend_date,
            "old_version_date": old_v["date"], "new_version_date": new_v["date"],
            "amendment_dates_in_range": [], "note": note, "changed_rows": changed_rows,
            "baseline_missing": baseline_missing,
        }

    old = fetch_version_articles(pcode, old_v)
    current = fetch_version_articles(pcode, new_v)
    all_nos = list(dict.fromkeys(list(old.keys()) + list(current.keys())))

    rows = []
    for no in all_nos:
        old_text, new_text = old.get(no), current.get(no)
        if old_text == new_text:
            continue
        if old_text is None:
            status, old_runs, new_runs = "新增", [], [(new_text, True)]
        elif new_text is None:
            status, old_runs, new_runs = "刪除", [(old_text, True)], []
        else:
            status = "修正"
            old_runs, new_runs = inline_diff(old_text, new_text)
        rows.append({"article_no": no, "status": status, "old_runs": old_runs, "new_runs": new_runs})

    amendment_dates_in_range = [v["date"] for v in reversed(versions) if old_v["date"] < v["date"] <= new_v["date"]]

    return {
        "pcode": pcode, "law_name": law_name, "amend_date": amend_date,
        "old_version_date": old_v["date"], "new_version_date": new_v["date"],
        "amendment_dates_in_range": amendment_dates_in_range, "note": note, "changed_rows": rows,
        "baseline_missing": baseline_missing,
    }


# ---------------------------------------------------------------------------
# 資料庫
# ---------------------------------------------------------------------------

APP_DIR = os.path.dirname(os.path.abspath(__file__))
# 本機執行時，資料庫預設放在程式旁邊；部署到雲端主機（例如 Railway）時，
# 用環境變數 LAW_TOOL_DB_PATH 指到「永久儲存空間（Volume）」掛載的路徑，
# 這樣服務重啟、重新部署時，已追蹤的法規和鑑別紀錄才不會被清空。
DB_PATH = os.environ.get("LAW_TOOL_DB_PATH") or os.path.join(APP_DIR, "law_tool.db")
os.makedirs(os.path.dirname(DB_PATH) or ".", exist_ok=True)


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():
    conn = get_db()
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS laws (
            pcode TEXT PRIMARY KEY,
            law_name TEXT,
            amend_date_display TEXT,
            note TEXT DEFAULT '',
            last_checked_until TEXT,
            created_at TEXT
        );
        CREATE TABLE IF NOT EXISTS findings (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            pcode TEXT NOT NULL,
            article_no TEXT NOT NULL,
            status TEXT,
            old_text TEXT,
            new_text TEXT,
            old_runs_json TEXT,
            new_runs_json TEXT,
            identified_date TEXT,
            applicability TEXT DEFAULT '',
            compliance_eval TEXT DEFAULT '',
            identifier TEXT DEFAULT '',
            remark TEXT DEFAULT '',
            closed INTEGER DEFAULT 0,
            UNIQUE(pcode, article_no)
        );
        """
    )
    conn.commit()
    conn.close()


def upsert_findings(pcode: str, changed_rows: list):
    """把這次比對到的異動條文寫進資料庫。已經存在、內容沒有再變的條文不動
    （保留你之前填好的評估結果）；內容又變了的條文，才重設評估欄位、
    提醒你重新鑑別。"""
    conn = get_db()
    today = date.today().isoformat()
    for row in changed_rows:
        old_text = "".join(t for t, _ in row["old_runs"])
        new_text = "".join(t for t, _ in row["new_runs"])
        existing = conn.execute(
            "SELECT new_text FROM findings WHERE pcode=? AND article_no=?",
            (pcode, row["article_no"]),
        ).fetchone()
        if existing and existing["new_text"] == new_text:
            continue  # 這條的最新內容跟資料庫裡記錄的一樣，不用重新鑑別
        if existing:
            conn.execute(
                """UPDATE findings SET status=?, old_text=?, new_text=?, old_runs_json=?,
                   new_runs_json=?, identified_date=?, applicability='', compliance_eval='',
                   closed=0 WHERE pcode=? AND article_no=?""",
                (
                    row["status"], old_text, new_text,
                    json.dumps(row["old_runs"], ensure_ascii=False),
                    json.dumps(row["new_runs"], ensure_ascii=False),
                    today, pcode, row["article_no"],
                ),
            )
        else:
            conn.execute(
                """INSERT INTO findings
                   (pcode, article_no, status, old_text, new_text, old_runs_json, new_runs_json,
                    identified_date)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                (
                    pcode, row["article_no"], row["status"], old_text, new_text,
                    json.dumps(row["old_runs"], ensure_ascii=False),
                    json.dumps(row["new_runs"], ensure_ascii=False),
                    today,
                ),
            )
    conn.commit()
    conn.close()


def runs_to_html(runs_json: str) -> Markup:
    runs = json.loads(runs_json) if runs_json else []
    if not runs:
        return Markup('<span class="muted">（本次修正未變動此條）</span>')
    parts = []
    for text, changed in runs:
        t = str(escape(text)).replace("\n", "<br>")
        parts.append(f"<b class='chg'>{t}</b>" if changed else t)
    return Markup("".join(parts))


# ---------------------------------------------------------------------------
# Flask App
# ---------------------------------------------------------------------------

app = Flask(__name__)
# secret_key 只用來簽署 flash 訊息用的 session cookie，這個工具沒有登入機制、
# 不存放帳密，風險很低；部署到雲端主機時可用環境變數 SECRET_KEY 覆寫。
app.secret_key = os.environ.get("SECRET_KEY", "law-tool-local-only")

BASE_CSS = """
:root{
  --ink:#202a24; --paper:#f3f2ec; --surface:#ffffff; --line:#dcd9cd;
  --muted:#6e6a5c; --accent:#a8341f; --accent-2:#3f6b4a;
}
*{box-sizing:border-box;}
body{margin:0;background:var(--paper);color:var(--ink);
  font-family:"Noto Sans TC","Microsoft JhengHei",system-ui,sans-serif;line-height:1.6;}
header{background:var(--surface);border-bottom:1px solid var(--line);padding:16px 28px;
  display:flex;align-items:center;justify-content:space-between;flex-wrap:wrap;gap:10px;}
header h1{font-size:18px;margin:0;}
header h1 a{color:var(--ink);text-decoration:none;}
main{max-width:1100px;margin:0 auto;padding:24px 20px 80px;}
.flash{background:#fdf1ee;border:1px solid var(--accent);color:var(--accent);
  padding:10px 14px;border-radius:6px;margin-bottom:16px;font-size:14px;}
table{border-collapse:collapse;width:100%;background:var(--surface);
  border:1px solid var(--line);border-radius:8px;overflow:hidden;font-size:13.5px;}
th,td{padding:10px 12px;border-bottom:1px solid var(--line);text-align:left;vertical-align:top;}
th{background:#faf9f4;font-weight:600;font-size:12.5px;color:var(--muted);}
tr:last-child td{border-bottom:none;}
.btn{display:inline-block;padding:7px 14px;border-radius:6px;border:1px solid var(--line);
  background:var(--surface);color:var(--ink);cursor:pointer;font-size:13.5px;text-decoration:none;}
.btn:hover{border-color:var(--ink);}
.btn-primary{background:var(--ink);color:var(--surface);border-color:var(--ink);}
.btn-danger{color:var(--accent);border-color:var(--accent);}
.card{background:var(--surface);border:1px solid var(--line);border-radius:8px;
  padding:20px;margin-bottom:20px;}
input[type=text],input[type=date],select{
  padding:6px 8px;border:1px solid var(--line);border-radius:5px;font-size:13.5px;
  font-family:inherit;width:100%;}
label{font-size:12.5px;color:var(--muted);display:block;margin-bottom:4px;}
.row{display:flex;gap:14px;flex-wrap:wrap;align-items:end;}
.row > div{flex:1;min-width:140px;}
.muted{color:var(--muted);}
.chg{text-decoration:underline;text-decoration-thickness:2px;}
.pill{display:inline-block;padding:2px 8px;border-radius:10px;font-size:11.5px;font-weight:600;}
.pill-open{background:#fdf1ee;color:var(--accent);}
.pill-closed{background:#eef4ef;color:var(--accent-2);}
.diff-cell{max-width:280px;}
.top-actions{display:flex;gap:10px;margin-bottom:18px;flex-wrap:wrap;}

/* 手機／窄螢幕 RWD：表格改成可以左右滑動，避免欄位被硬擠爆版；
   按鈕、輸入框加大，方便手指點按；整體邊距縮小，多留一點內容空間。 */
.table-scroll{overflow-x:auto;-webkit-overflow-scrolling:touch;border-radius:8px;}
.table-scroll table{border-radius:0;}
@media (max-width:700px){
  main{padding:16px 14px 60px;}
  header{padding:12px 16px;}
  header h1{font-size:16px;}
  .card{padding:16px;}
  .table-scroll table{min-width:640px;}
  .table-scroll table.table-wide{min-width:980px;}
  .btn{padding:9px 14px;}
  input[type=text],input[type=date],select{font-size:16px;padding:8px;}
  .row{gap:10px;}
}
"""

LAYOUT = """
<!doctype html>
<html lang="zh-Hant">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{{ title or "法規鑑別追蹤工具" }}</title>
<style>{{ css|safe }}</style>
</head>
<body>
<header>
  <h1><a href="{{ url_for('dashboard') }}">法規鑑別追蹤工具</a></h1>
  <div class="muted" style="font-size:12.5px;">資料來源：勞動部勞動法令查詢系統（laws.mol.gov.tw）</div>
</header>
<main>
  {% with messages = get_flashed_messages() %}
    {% for m in messages %}<div class="flash">{{ m }}</div>{% endfor %}
  {% endwith %}
  {{ body|safe }}
</main>
</body>
</html>
"""


def render(body_html, **ctx):
    body = render_template_string(body_html, **ctx)
    return render_template_string(LAYOUT, css=BASE_CSS, body=body, **ctx)


DASHBOARD_BODY = """
<div class="top-actions">
  <form method="post" action="{{ url_for('check_all') }}">
    <button class="btn btn-primary" type="submit">全部立即比對</button>
  </form>
  <a class="btn" href="{{ url_for('export_all') }}">匯出全部法規 Excel</a>
  <form method="post" action="{{ url_for('migrate_pcode') }}"
        onsubmit="return confirm('這是資料來源改成勞動法令查詢系統後的一次性修復：會嘗試把清單裡還在用舊代碼的法規自動換成新代碼，既有的鑑別紀錄不會不見。確定要執行嗎？');">
    <button class="btn" type="submit">修復舊資料來源代碼</button>
  </form>
</div>

<div class="card">
  <h3 style="margin-top:0;">新增要追蹤的法規</h3>
  <form method="get" action="{{ url_for('search_laws') }}" style="margin-bottom:14px;">
    <div class="row">
      <div>
        <label>用法規名稱關鍵字搜尋（推薦，不用打完整全名，例如「教育訓練規則」）</label>
        <input type="text" name="kw" placeholder="輸入法規名稱關鍵字">
      </div>
      <div style="flex:0;min-width:150px;">
        <label>追蹤起始日（套用到搜尋結果裡勾選的法規）</label>
        <input type="date" name="start_date" value="{{ year_start }}">
      </div>
      <div style="flex:0;min-width:150px;">
        <label>追蹤結束日（選填）</label>
        <input type="date" name="end_date" value="{{ today }}">
      </div>
      <div style="flex:0;">
        <button class="btn btn-primary" type="submit">搜尋法規</button>
      </div>
    </div>
  </form>
  <details>
    <summary class="muted" style="cursor:pointer;font-size:12.5px;">或者，直接輸入法規代碼（pcode）新增</summary>
    <form method="post" action="{{ url_for('add_law') }}" style="margin-top:12px;">
      <div class="row">
        <div>
          <label>法規代碼（可從 laws.mol.gov.tw 該法規網址的 id= 後面找到）</label>
          <input type="text" name="pcode" placeholder="例如 FL015013" required>
        </div>
        <div>
          <label>追蹤起始日（只比對這天之後發生的修正）</label>
          <input type="date" name="start_date" value="{{ year_start }}" required>
        </div>
        <div style="flex:0;">
          <button class="btn btn-primary" type="submit">新增</button>
        </div>
      </div>
    </form>
  </details>
</div>

<div class="row" style="align-items:end;margin-bottom:10px;">
  <div style="flex:0;min-width:230px;">
    <label>排序方式</label>
    <select onchange="location.href='{{ url_for('dashboard') }}?sort='+this.value">
      <option value="" {{ "selected" if not sort else "" }}>預設（新增順序）</option>
      <option value="amend_desc" {{ "selected" if sort=="amend_desc" else "" }}>現行修正日期：新→舊</option>
      <option value="amend_asc" {{ "selected" if sort=="amend_asc" else "" }}>現行修正日期：舊→新</option>
      <option value="checked_desc" {{ "selected" if sort=="checked_desc" else "" }}>已追蹤到：新→舊</option>
      <option value="checked_asc" {{ "selected" if sort=="checked_asc" else "" }}>已追蹤到：舊→新</option>
      <option value="open_first" {{ "selected" if sort=="open_first" else "" }}>待鑑別：未結案優先</option>
      <option value="open_last" {{ "selected" if sort=="open_last" else "" }}>待鑑別：已結案優先</option>
    </select>
  </div>
</div>

<div class="table-scroll">
<table>
  <thead><tr>
    <th>法規名稱</th><th>pcode</th><th>現行修正日期</th><th>已追蹤到</th>
    <th>待鑑別</th><th>操作</th>
  </tr></thead>
  <tbody>
    {% for law in laws %}
    <tr>
      <td><a href="{{ url_for('law_detail', pcode=law.pcode) }}">{{ law.law_name or law.pcode }}</a>
        {% if law.note %}<div style="font-size:12px;color:var(--accent);">⚠ {{ law.note }}</div>{% endif %}
      </td>
      <td class="muted">{{ law.pcode }}</td>
      <td>{{ law.amend_date_display or "-" }}</td>
      <td>{{ law.last_checked_display }}</td>
      <td>
        {% if law.open_count %}
        <span class="pill pill-open">{{ law.open_count }} 筆未結案</span>
        {% else %}
        <span class="pill pill-closed">全部結案</span>
        {% endif %}
      </td>
      <td style="white-space:nowrap;">
        <form style="display:inline" method="post" action="{{ url_for('check_law', pcode=law.pcode) }}">
          <button class="btn" type="submit">立即比對</button>
        </form>
        <a class="btn" href="{{ url_for('law_detail', pcode=law.pcode) }}">查看</a>
        <form style="display:inline" method="post" action="{{ url_for('delete_law', pcode=law.pcode) }}"
              onsubmit="return confirm('確定要刪除這部法規跟它累積的鑑別紀錄嗎？');">
          <button class="btn btn-danger" type="submit">刪除</button>
        </form>
      </td>
    </tr>
    {% else %}
    <tr><td colspan="6" class="muted">還沒有追蹤任何法規，先在上面新增一部看看。</td></tr>
    {% endfor %}
  </tbody>
</table>
</div>
"""


def _sort_dashboard_rows(rows: list, sort: str) -> list:
    """依照下拉選單選的排序方式排序法規清單；沒給或給不認得的值就維持原本
    （新增順序）不動。日期缺值一律排到最後，不管是新→舊還是舊→新。"""
    if sort == "amend_desc":
        rows.sort(key=lambda d: roc_to_yyyymmdd(d["amend_date_display"]) or "", reverse=True)
    elif sort == "amend_asc":
        rows.sort(key=lambda d: roc_to_yyyymmdd(d["amend_date_display"]) or "99999999")
    elif sort == "checked_desc":
        rows.sort(key=lambda d: d["last_checked_until"] or "", reverse=True)
    elif sort == "checked_asc":
        rows.sort(key=lambda d: d["last_checked_until"] or "99999999")
    elif sort == "open_first":
        rows.sort(key=lambda d: -d["open_count"])
    elif sort == "open_last":
        rows.sort(key=lambda d: d["open_count"])
    return rows


def _migrate_stale_pcodes() -> dict:
    """一次性搬家用：資料來源從全國法規資料庫換成勞動法令查詢系統之後，原本
    存在資料庫裡的舊代碼（例如 N0060001）在新資料來源查不到對應資料，直接
    比對會失敗。這裡逐筆檢查現有的追蹤清單：先直接拿現有代碼去問新資料
    來源，抓得到法規名稱就代表這筆代碼本來就沒事、不用動；抓不到的話改用
    法規名稱去新資料來源搜尋一次，如果剛好只找到一筆同名結果就自動把
    laws、findings 兩張表裡的代碼都換成新的（換代碼不會動到既有的鑑別紀錄、
    備註、已比對到哪一天，因為這些欄位本來就是用代碼去關聯，代碼換了資料
    還在）；找不到或找到好幾筆同名結果（沒辦法自動判斷哪一筆才對應）就跳過，
    回報請使用者自行到「新增要追蹤的法規」用關鍵字搜尋、手動確認後重新
    整理。"""
    conn = get_db()
    laws = conn.execute("SELECT * FROM laws").fetchall()
    migrated, failed, unchanged = [], [], []
    for law in laws:
        pcode, law_name = law["pcode"], law["law_name"]
        try:
            info = list_versions(pcode)
            if info["law_name"]:
                unchanged.append(law_name)
                continue
        except Exception:
            pass
        try:
            candidates, _truncated = search_laws_by_keyword(law_name)
        except Exception as e:
            failed.append(f"{law_name}（搜尋失敗：{e}）")
            continue
        exact = [r for r in candidates if r["law_name"] == law_name]
        matches = exact or candidates
        if len(matches) != 1:
            failed.append(f"{law_name}（自動搜尋找到 {len(matches)} 筆可能對應的結果，需自行確認）")
            continue
        new_pcode = matches[0]["pcode"]
        if new_pcode == pcode:
            unchanged.append(law_name)
            continue
        try:
            conn.execute("UPDATE laws SET pcode=? WHERE pcode=?", (new_pcode, pcode))
            conn.execute("UPDATE OR IGNORE findings SET pcode=? WHERE pcode=?", (new_pcode, pcode))
        except sqlite3.IntegrityError as e:
            failed.append(f"{law_name}（新代碼 {new_pcode} 更新失敗：{e}）")
            continue
        migrated.append(f"{law_name}（{pcode} → {new_pcode}）")
    conn.commit()
    conn.close()
    return {"migrated": migrated, "failed": failed, "unchanged": unchanged}


@app.route("/migrate_pcode", methods=["POST"])
def migrate_pcode():
    result = _migrate_stale_pcodes()
    if result["migrated"]:
        flash("已自動更新為新資料來源的代碼：" + "、".join(result["migrated"]))
    if result["failed"]:
        flash("以下法規需要自行手動處理（建議刪除後用關鍵字搜尋重新新增）：" + "、".join(result["failed"]))
    if not result["migrated"] and not result["failed"]:
        flash("檢查完成，所有法規的代碼都已經可以在新資料來源上正常使用，不需要更新。")
    return redirect(url_for("dashboard"))


@app.route("/")
def dashboard():
    sort = request.args.get("sort", "")
    conn = get_db()
    laws = conn.execute("SELECT * FROM laws ORDER BY created_at").fetchall()
    rows = []
    for law in laws:
        open_count = conn.execute(
            "SELECT COUNT(*) c FROM findings WHERE pcode=? AND closed=0", (law["pcode"],)
        ).fetchone()["c"]
        d = dict(law)
        d["open_count"] = open_count
        d["last_checked_display"] = fmt_date(law["last_checked_until"]) if law["last_checked_until"] else "尚未比對"
        rows.append(d)
    conn.close()
    rows = _sort_dashboard_rows(rows, sort)
    year_start = date(date.today().year, 1, 1).isoformat()
    return render(DASHBOARD_BODY, laws=rows, today=date.today().isoformat(), year_start=year_start, sort=sort)


def _add_law_by_pcode(pcode: str, start_date: str, end_date: str = None) -> tuple:
    """新增一部要追蹤的法規（給單筆新增、批次新增共用）。成功回傳
    (法規名稱, 比對訊息或 None)；失敗時丟出 requests.RequestException 或
    RuntimeError，由呼叫端決定怎麼顯示。

    有給 end_date 的話，代表使用者選的不只是「追蹤起始日」，還有明確的
    「結束日期」——這種情況新增完馬上就跑一次比對（範圍剛好是
    起始日～結束日），直接把這段期間內的異動存進待鑑別清單，不用再等她
    自己另外按「立即比對」；沒給 end_date（例如舊的單筆輸入 pcode 表單）
    就維持原本的行為，只設定追蹤起始日當基準，之後再由使用者自己按
    「立即比對」。"""
    info = list_versions(pcode)
    if not info["law_name"]:
        raise RuntimeError(f"抓不到 pcode={pcode} 的法規名稱，請確認代碼是否正確。")

    # last_checked_until 的定義統一是「已經確認到這一天為止，之後的檢查只看
    # 更新的部分」；所以「追蹤起始日」要換算成「起始日前一天」存起來，這樣第一次
    # 「立即比對」查詢時（since = next_day(last_checked_until)）才會剛好從
    # 使用者選的那一天開始算，不會漏掉剛好在起始日當天生效的修正。
    baseline = previous_day(start_date) if start_date else None

    conn = get_db()
    conn.execute(
        "INSERT OR REPLACE INTO laws (pcode, law_name, amend_date_display, note, last_checked_until, created_at) "
        "VALUES (?, ?, ?, COALESCE((SELECT note FROM laws WHERE pcode=?), ''), ?, "
        "COALESCE((SELECT created_at FROM laws WHERE pcode=?), ?))",
        (pcode, info["law_name"], info["amend_date_display"], pcode, baseline, pcode, datetime.now().isoformat()),
    )
    conn.commit()
    conn.close()

    check_msg = _run_check(pcode, until=end_date) if end_date else None
    return info["law_name"], check_msg


@app.route("/laws", methods=["POST"])
def add_law():
    pcode = request.form.get("pcode", "").strip()
    start_date = request.form.get("start_date", "").strip().replace("-", "")
    end_date = request.form.get("end_date", "").strip().replace("-", "")
    if not pcode:
        flash("請輸入法規代碼。")
        return redirect(url_for("dashboard"))
    try:
        law_name, check_msg = _add_law_by_pcode(pcode, start_date, end_date or None)
    except requests.RequestException as e:
        flash(f"連線失敗：{e}")
        return redirect(url_for("dashboard"))
    except RuntimeError as e:
        flash(str(e))
        return redirect(url_for("dashboard"))
    msg = f"已新增「{law_name}」，追蹤起始日設為 {fmt_date(start_date)}。"
    if check_msg:
        msg += f" {check_msg}"
    flash(msg)
    return redirect(url_for("dashboard"))


@app.route("/laws/batch", methods=["POST"])
def add_laws_batch():
    pcodes = [p.strip() for p in request.form.getlist("selected_pcodes") if p.strip()]
    # 這兩個表單欄位是 <input type="date">，瀏覽器送出的格式是「YYYY-MM-DD」；
    # 網址查詢字串（給 search_laws 重新導向用）要保留這個帶槓的格式，日期
    # 輸入框才讀得回去，內部比對邏輯用的「YYYYMMDD」緊湊格式則另外去掉槓
    # 再處理，兩種格式不要弄混。
    start_date_iso = request.form.get("start_date", "").strip()
    end_date_iso = request.form.get("end_date", "").strip()
    start_date = start_date_iso.replace("-", "")
    end_date = end_date_iso.replace("-", "")
    keyword = request.form.get("kw", "")
    sort = request.form.get("sort", "")

    if not pcodes:
        flash("請至少勾選一部法規再按追蹤。")
        return redirect(url_for(
            "search_laws", kw=keyword, sort=sort, start_date=start_date_iso, end_date=end_date_iso
        ))

    added, failed, check_msgs = [], [], []
    for pcode in pcodes:
        try:
            law_name, check_msg = _add_law_by_pcode(pcode, start_date, end_date or None)
            added.append(law_name)
            if check_msg:
                check_msgs.append(check_msg)
        except (requests.RequestException, RuntimeError) as e:
            failed.append(f"{pcode}（{e}）")

    msgs = []
    if added:
        range_text = f"{fmt_date(start_date)}"
        if end_date:
            range_text += f" ～ {fmt_date(end_date)}"
        msgs.append(f"已新增追蹤 {len(added)} 部法規，範圍 {range_text}：" + "、".join(added))
    msgs.extend(check_msgs)
    if failed:
        msgs.append(f"{len(failed)} 部新增失敗：" + "、".join(failed))
    flash(" ／ ".join(msgs))
    return redirect(url_for("dashboard"))


SEARCH_BODY = """
<p><a href="{{ url_for('dashboard') }}">&larr; 回法規清單</a></p>
<div class="card">
  <h3 style="margin-top:0;">用法規名稱關鍵字搜尋</h3>
  <form method="get" action="{{ url_for('search_laws') }}">
    <div class="row">
      <div>
        <label>法規名稱關鍵字（不用打完整全名，例如「教育訓練規則」）</label>
        <input type="text" name="kw" value="{{ keyword }}" placeholder="例如：教育訓練規則" autofocus>
      </div>
      <div style="flex:0;min-width:150px;">
        <label>追蹤起始日（套用到下面所有勾選的法規）</label>
        <input type="date" name="start_date" value="{{ start_date }}">
      </div>
      <div style="flex:0;min-width:150px;">
        <label>追蹤結束日（選填，不填就先只設定起始日）</label>
        <input type="date" name="end_date" value="{{ end_date }}">
      </div>
      <input type="hidden" name="sort" value="{{ sort }}">
      <div style="flex:0;">
        <button class="btn btn-primary" type="submit">搜尋</button>
      </div>
    </div>
  </form>
</div>

{% if error %}
<div class="flash">{{ error }}</div>
{% elif keyword %}
  {% if results %}
  <form method="post" action="{{ url_for('add_laws_batch') }}">
  <input type="hidden" name="kw" value="{{ keyword }}">
  <input type="hidden" name="start_date" value="{{ start_date }}">
  <input type="hidden" name="end_date" value="{{ end_date }}">
  <div class="card" style="padding:14px 16px;margin-bottom:12px;">
    <div class="row" style="align-items:end;">
      <div class="muted" style="font-size:12.5px;">
        追蹤範圍：{{ start_date }}{% if end_date %} ～ {{ end_date }}{% endif %}
        （要改的話調整上面搜尋列的日期後重新搜尋一次）
      </div>
      <div style="flex:0;">
        <button class="btn btn-primary" type="submit">追蹤已勾選的法規</button>
      </div>
      <div class="muted" style="font-size:12.5px;">先勾選左邊要追蹤的法規（可複選），再按這裡一次新增</div>
    </div>
  </div>
  <div class="row" style="align-items:end;margin-bottom:10px;">
    <div style="flex:0;min-width:200px;">
      <label>排序方式</label>
      <select onchange="location.href='{{ url_for('search_laws', kw=keyword, start_date=start_date, end_date=end_date) }}&sort='+this.value">
        <option value="" {{ "selected" if not sort else "" }}>預設（符合度）</option>
        <option value="amend_desc" {{ "selected" if sort=="amend_desc" else "" }}>現行修正日期：新→舊</option>
        <option value="amend_asc" {{ "selected" if sort=="amend_asc" else "" }}>現行修正日期：舊→新</option>
      </select>
    </div>
  </div>
  <div class="table-scroll">
  <table>
    <thead><tr>
      <th style="width:40px;">
        <input type="checkbox"
          onclick="var checked=this.checked;this.closest('table').querySelectorAll('.pick-law').forEach(function(c){c.checked=checked;});"
          title="全選／全不選">
      </th>
      <th>法規名稱</th><th>現行修正日期</th><th>pcode</th><th style="width:140px;">狀態</th>
    </tr></thead>
    <tbody>
      {% for r in results %}
      <tr>
        <td>
          {% if r.pcode in tracked %}
          <input type="checkbox" disabled title="已在追蹤清單">
          {% else %}
          <input type="checkbox" class="pick-law" name="selected_pcodes" value="{{ r.pcode }}">
          {% endif %}
        </td>
        <td>{{ r.law_name }}{% if r.repealed %} <span class="pill pill-open">已廢止</span>{% endif %}</td>
        <td>{{ r.amend_date_display or "-" }}</td>
        <td class="muted">{{ r.pcode }}</td>
        <td>
          {% if r.pcode in tracked %}
          <a class="btn" href="{{ url_for('law_detail', pcode=r.pcode) }}">已在追蹤，查看</a>
          {% else %}
          <span class="muted">未追蹤</span>
          {% endif %}
        </td>
      </tr>
      {% endfor %}
    </tbody>
  </table>
  </div>
  </form>
  {% if truncated %}
  <p class="muted" style="margin-top:10px;">符合的結果較多，這裡只列出前面幾筆，如果沒看到你要的法規，
    請輸入更精確的關鍵字（例如加上「辦法」「規則」「標準」等法規名稱裡的字）。</p>
  {% endif %}
  {% else %}
  <p class="muted">找不到名稱包含「{{ keyword }}」的法規，換個關鍵字試試看
    （例如只打法規名稱裡最特別的幾個字，不用打「XX法」「XX辦法」這種通用字）。</p>
  {% endif %}
{% endif %}
"""


@app.route("/search")
def search_laws():
    keyword = request.args.get("kw", "").strip()
    sort = request.args.get("sort", "")
    today_str = date.today().isoformat()
    year_start_str = date(date.today().year, 1, 1).isoformat()
    # 追蹤起始日／結束日跟著查詢字串走：第一次進來（網址上沒帶這兩個參數）
    # 才套用預設值（起始日＝今年 1/1，結束日＝今天），只要使用者改過、重新
    # 搜尋過一次，這兩個值就會原封不動地跟著新的搜尋結果一起帶回來，不會
    # 搜尋完又跳回預設值，方便直接接著勾選、追蹤。
    start_date = request.args.get("start_date", "").strip() or year_start_str
    end_date = request.args.get("end_date", "").strip() or today_str
    results, truncated, error = [], False, None
    if keyword:
        try:
            results, truncated = search_laws_by_keyword(keyword)
        except requests.RequestException as e:
            error = f"連線失敗：{e}"
    if sort == "amend_desc":
        results.sort(key=lambda r: roc_to_yyyymmdd(r["amend_date_display"]) or "", reverse=True)
    elif sort == "amend_asc":
        results.sort(key=lambda r: roc_to_yyyymmdd(r["amend_date_display"]) or "99999999")
    conn = get_db()
    tracked = {r["pcode"] for r in conn.execute("SELECT pcode FROM laws").fetchall()}
    conn.close()
    return render(
        SEARCH_BODY, keyword=keyword, results=results, truncated=truncated,
        error=error, tracked=tracked, today=today_str, sort=sort,
        start_date=start_date, end_date=end_date,
    )


@app.route("/laws/<pcode>/delete", methods=["POST"])
def delete_law(pcode):
    conn = get_db()
    conn.execute("DELETE FROM findings WHERE pcode=?", (pcode,))
    conn.execute("DELETE FROM laws WHERE pcode=?", (pcode,))
    conn.commit()
    conn.close()
    flash("已刪除。")
    return redirect(url_for("dashboard"))


def _run_check(pcode: str, until: str = None):
    """跑一次比對；until 預設 None 表示比對到今天（平常「立即比對」用的行為），
    新增法規時如果有給明確的「結束日期」，會帶著 until 呼叫這裡，把第一次
    比對限定在使用者選的「起始日～結束日」這段範圍內，而不是自動比到今天。"""
    conn = get_db()
    law = conn.execute("SELECT * FROM laws WHERE pcode=?", (pcode,)).fetchone()
    conn.close()
    if not law:
        return f"[{pcode}] 找不到這部法規。"
    # last_checked_until 存的是「已經確認到這一天為止」，所以查詢時要從
    # 它的隔天開始算，才不會把已經看過、已經填好評估的那次異動又抓回來一次。
    since = next_day(law["last_checked_until"]) if law["last_checked_until"] else None
    try:
        report = build_report(pcode, since=since, until=until)
    except (requests.RequestException, RuntimeError) as e:
        return f"[{report_name(law)}] 比對失敗：{e}"

    upsert_findings(pcode, report["changed_rows"])

    new_baseline = report["new_version_date"] or law["last_checked_until"] or date.today().strftime("%Y%m%d")
    # 把 build_report 算出來的 note（例如「資料庫最早只能追溯到 XXXX」這種
    # 追溯範圍受限的提醒）存回法規清單，讓首頁那一列也看得到，不會像以前
    # 一樣算出來卻沒地方顯示、直接被吞掉。之後追蹤基準往前推進、不再有這個
    # 限制時，note 就會是空字串，畫面上自然不會再顯示。
    conn = get_db()
    conn.execute(
        "UPDATE laws SET last_checked_until=?, law_name=?, amend_date_display=?, note=? WHERE pcode=?",
        (new_baseline, report["law_name"], report["amend_date"], report.get("note") or "", pcode),
    )
    conn.commit()
    conn.close()

    n = len(report["changed_rows"])
    if n == 0:
        return f"「{report['law_name']}」：這段期間沒有新的修正。"
    if report.get("baseline_missing"):
        # baseline_missing 代表這幾筆是靠沿革公告文字解析出來、不是逐字比對舊
        # 條文得到的結果，訊息要講清楚、跟真的比對到差異的情況分開，提醒她
        # 「修正」「刪除」的條文還是要自己核對過才能結案。
        return (
            f"「{report['law_name']}」：⚠ 追蹤期間內有修正，但資料庫沒有保留舊條文可自動比對，"
            f"已根據沿革公告解析出 {n} 條可能異動的條文存入待鑑別清單，請自行核對後手動鑑別、結案。"
        )
    return f"「{report['law_name']}」：新增／更新了 {n} 條異動，已存入待鑑別清單。"


def report_name(law_row):
    return law_row["law_name"] or law_row["pcode"]


@app.route("/laws/<pcode>/check", methods=["POST"])
def check_law(pcode):
    msg = _run_check(pcode)
    flash(msg)
    return redirect(request.referrer or url_for("dashboard"))


@app.route("/check_all", methods=["POST"])
def check_all():
    conn = get_db()
    pcodes = [r["pcode"] for r in conn.execute("SELECT pcode FROM laws").fetchall()]
    conn.close()
    msgs = [_run_check(p) for p in pcodes]
    flash(" ／ ".join(msgs) if msgs else "還沒有追蹤任何法規。")
    return redirect(url_for("dashboard"))


LAW_DETAIL_BODY = """
<p><a href="{{ url_for('dashboard') }}">&larr; 回法規清單</a></p>
<div class="card">
  <h2 style="margin:0 0 4px;">{{ law.law_name }}</h2>
  <div class="muted">pcode={{ law.pcode }}　｜　已追蹤到：{{ last_checked_display }}</div>
</div>

<div class="top-actions">
  <form method="post" action="{{ url_for('check_law', pcode=law.pcode) }}">
    <button class="btn btn-primary" type="submit">立即比對</button>
  </form>
  <a class="btn" href="{{ url_for('export_one', pcode=law.pcode) }}">匯出此法規 Excel</a>
</div>

<form method="post" action="{{ url_for('save_findings', pcode=law.pcode) }}">
<div class="table-scroll">
<table class="table-wide">
  <thead><tr>
    <th style="width:5%">條號</th>
    <th style="width:8%">異動</th>
    <th style="width:20%">修正前</th>
    <th style="width:20%">修正後</th>
    <th style="width:10%">鑑別日期</th>
    <th style="width:9%">條文適用性</th>
    <th style="width:14%">守規性之評估</th>
    <th style="width:8%">鑑別人員</th>
    <th style="width:10%">備註</th>
    <th style="width:6%">結案</th>
  </tr></thead>
  <tbody>
    {% for f in findings %}
    <tr>
      <td>{{ f.article_no }}</td>
      <td>{{ f.status }}</td>
      <td class="diff-cell">{{ f.old_html }}</td>
      <td class="diff-cell">{{ f.new_html }}</td>
      <td class="muted">{{ f.identified_date }}</td>
      <td>
        <select name="applicability_{{ f.id }}">
          {% for opt in ["", "適用", "參考", "不適用"] %}
          <option value="{{ opt }}" {{ "selected" if f.applicability==opt else "" }}>{{ opt or "（未選）" }}</option>
          {% endfor %}
        </select>
      </td>
      <td><input type="text" name="compliance_eval_{{ f.id }}" value="{{ f.compliance_eval }}"></td>
      <td><input type="text" name="identifier_{{ f.id }}" value="{{ f.identifier }}"></td>
      <td><input type="text" name="remark_{{ f.id }}" value="{{ f.remark }}"></td>
      <td style="text-align:center;">
        <input type="checkbox" name="closed_{{ f.id }}" {{ "checked" if f.closed else "" }}>
      </td>
    </tr>
    {% else %}
    <tr><td colspan="10" class="muted">目前沒有待鑑別的條文，按上面「立即比對」試試看。</td></tr>
    {% endfor %}
  </tbody>
</table>
</div>
{% if findings %}
<p style="margin-top:14px;"><button class="btn btn-primary" type="submit">儲存這一頁的填寫內容</button></p>
{% endif %}
</form>
"""


@app.route("/laws/<pcode>")
def law_detail(pcode):
    conn = get_db()
    law = conn.execute("SELECT * FROM laws WHERE pcode=?", (pcode,)).fetchone()
    if not law:
        conn.close()
        flash("找不到這部法規。")
        return redirect(url_for("dashboard"))
    findings = conn.execute(
        "SELECT * FROM findings WHERE pcode=? ORDER BY closed, identified_date DESC, id", (pcode,)
    ).fetchall()
    conn.close()

    rows = []
    for f in findings:
        d = dict(f)
        d["old_html"] = runs_to_html(f["old_runs_json"])
        d["new_html"] = runs_to_html(f["new_runs_json"])
        rows.append(d)

    last_checked_display = fmt_date(law["last_checked_until"]) if law["last_checked_until"] else "尚未比對"
    return render(LAW_DETAIL_BODY, law=law, findings=rows, last_checked_display=last_checked_display)


@app.route("/laws/<pcode>/save", methods=["POST"])
def save_findings(pcode):
    conn = get_db()
    ids = [r["id"] for r in conn.execute("SELECT id FROM findings WHERE pcode=?", (pcode,)).fetchall()]
    for fid in ids:
        applicability = request.form.get(f"applicability_{fid}", "")
        compliance_eval = request.form.get(f"compliance_eval_{fid}", "")
        identifier = request.form.get(f"identifier_{fid}", "")
        remark = request.form.get(f"remark_{fid}", "")
        closed = 1 if request.form.get(f"closed_{fid}") else 0
        conn.execute(
            """UPDATE findings SET applicability=?, compliance_eval=?, identifier=?, remark=?, closed=?
               WHERE id=?""",
            (applicability, compliance_eval, identifier, remark, closed, fid),
        )
    conn.commit()
    conn.close()
    flash("已儲存。")
    return redirect(url_for("law_detail", pcode=pcode))


# ---------------------------------------------------------------------------
# Excel 匯出
# ---------------------------------------------------------------------------


def safe_sheet_name(name: str, used: set) -> str:
    cleaned = re.sub(r"[\\/?*\[\]:]", "", name or "").strip() or "鑑別草稿"
    base = cleaned[:31]
    candidate = base
    n = 2
    while candidate in used:
        suffix = f"({n})"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate)
    return candidate


def write_law_sheet_from_db(wb, pcode: str, used_names: set):
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    from openpyxl.styles import Alignment, Font

    conn = get_db()
    law = conn.execute("SELECT * FROM laws WHERE pcode=?", (pcode,)).fetchone()
    findings = conn.execute(
        "SELECT * FROM findings WHERE pcode=? ORDER BY id", (pcode,)
    ).fetchall()
    conn.close()
    if not law or not findings:
        return False

    ws = wb.create_sheet(safe_sheet_name(law["law_name"], used_names))
    headers = [
        "序號", "法規名稱", "法規條文內容(修正前)", "法規條文內容(修正後)", "異動類型",
        "鑑別日期", "條文適用性\n(適用/參考/不適用)", "守規性之評估", "鑑別人員",
        "是否結案", "備註",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")

    bold_underline = InlineFont(b=True, u="single")
    article_label_font = InlineFont(b=True)

    def to_rich_text(article_no, runs_json):
        runs = json.loads(runs_json) if runs_json else []
        blocks = [TextBlock(article_label_font, f"{article_no}\n")]
        if not runs:
            blocks.append("（本次修正未變動此條）")
            return CellRichText(*blocks)
        for text, changed in runs:
            if not text:
                continue
            blocks.append(TextBlock(bold_underline, text) if changed else text)
        return CellRichText(*blocks)

    for i, f in enumerate(findings, start=1):
        ws.append(
            [
                i, law["law_name"],
                to_rich_text(f["article_no"], f["old_runs_json"]),
                to_rich_text(f["article_no"], f["new_runs_json"]),
                f["status"], f["identified_date"], f["applicability"], f["compliance_eval"],
                f["identifier"], "已結案" if f["closed"] else "未結案", f["remark"],
            ]
        )

    widths = [6, 20, 40, 40, 8, 12, 14, 20, 10, 8, 16]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = w
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return True


@app.route("/export/<pcode>")
def export_one(pcode):
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    ok = write_law_sheet_from_db(wb, pcode, set())
    if not ok:
        flash("這部法規目前沒有可匯出的鑑別紀錄。")
        return redirect(url_for("law_detail", pcode=pcode))
    out_path = os.path.join(APP_DIR, f"{pcode}_鑑別結果.xlsx")
    wb.save(out_path)
    return send_file(out_path, as_attachment=True)


@app.route("/export/all")
def export_all():
    from openpyxl import Workbook

    conn = get_db()
    pcodes = [r["pcode"] for r in conn.execute("SELECT pcode FROM laws").fetchall()]
    conn.close()

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    written = 0
    for pcode in pcodes:
        if write_law_sheet_from_db(wb, pcode, used):
            written += 1

    if written == 0:
        flash("目前沒有任何法規有可匯出的鑑別紀錄。")
        return redirect(url_for("dashboard"))

    out_path = os.path.join(APP_DIR, "全部法規_鑑別結果.xlsx")
    wb.save(out_path)
    return send_file(out_path, as_attachment=True)


# ---------------------------------------------------------------------------
# 啟動
# ---------------------------------------------------------------------------

# 不管是本機用 python law_tool_app.py 直接執行，還是部署到雲端主機用
# gunicorn 啟動（gunicorn 只會 import 這支檔案、不會執行下面的
# if __name__ == "__main__" 區塊），都要確保資料表存在，所以放在最外層。
init_db()


def _open_browser(url):
    time.sleep(1.0)
    webbrowser.open(url)


if __name__ == "__main__":
    # 本機雙擊 start_law_tool.bat／直接執行時走這條路：只監聽本機
    # 127.0.0.1，並自動開瀏覽器，跟以前的行為完全一樣。
    # （部署到雲端主機時，是用 gunicorn 啟動，不會跑到這裡，改由
    # gunicorn 監聽 0.0.0.0 上的 $PORT。）
    port = int(os.environ.get("PORT", 5050))
    url = f"http://127.0.0.1:{port}"
    threading.Thread(target=_open_browser, args=(url,), daemon=True).start()
    print(f"法規鑑別追蹤工具已啟動：{url}（要停止的話回到這個視窗按 Ctrl+C）")
    app.run(host="127.0.0.1", port=port, debug=False)
